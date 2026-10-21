"""Genera docs/backlog/*.md y docs/historias/*.md desde backlog_v2_4sprints.xlsx (fuente del curso)."""
import json
import re

import openpyxl

import pathlib
RAIZ = pathlib.Path(__file__).resolve().parents[2]
XLSX = str(RAIZ / 'docs/backlog/backlog_v2_4sprints.xlsx')
DOCS = str(RAIZ / 'docs')
wb = openpyxl.load_workbook(XLSX, data_only=True)

hist = []
for r in wb['Historias de Usuario'].iter_rows(min_row=5, max_row=24, min_col=2, max_col=11, values_only=True):
    if r[0]:
        hist.append(dict(id=r[0], enunciado=r[1], alias=r[2], estado=r[3] or '—', valor=r[4], esfuerzo=r[5],
                         roi=r[6], sprint=r[7], prioridad=r[8], comentario=r[9] or ''))
CA = {}
for r in wb['Criterios de aceptación'].iter_rows(min_row=5, values_only=True):
    if r[0]:
        CA.setdefault(r[0], []).append(dict(n=r[2], texto=r[3], prueba=r[4] or '', vinculo=r[5] or ''))
T = {}
for r in wb['Tareas (Sprint Backlog)'].iter_rows(min_row=5, values_only=True):
    if r[0]:
        T.setdefault(r[2], []).append(dict(id=r[0], sprint=r[1], tarea=r[3], tipo=r[4], resp=r[5], rev=r[6],
                                           horas=r[7], vinc=r[9] or ''))
REG = [r for r in wb['Reglas MinEdu'].iter_rows(min_row=5, max_row=13, max_col=4, values_only=True) if r[0]]
CAS = [r for r in wb['Reglas MinEdu'].iter_rows(min_row=17, max_row=21, max_col=12, values_only=True) if r[0]]
DD = wb['DoR y DoD']
dor = [r for r in DD.iter_rows(min_row=4, max_row=9, max_col=2, values_only=True) if r[0]]
dod = [r for r in DD.iter_rows(min_row=11, max_row=17, max_col=2, values_only=True) if r[0]]
HZ = [r for r in wb['Hallazgos'].iter_rows(min_row=4, max_col=5, values_only=True) if r[0]]


def tabla(head, rows):
    esc = lambda v: str('' if v is None else v).replace('|', '\\|').replace('\n', ' ')
    out = ['| ' + ' | '.join(head) + ' |', '|' + '|'.join(['---'] * len(head)) + '|']
    return '\n'.join(out + ['| ' + ' | '.join(esc(c) for c in row) + ' |' for row in rows])


def slug(s):
    for a, b in zip('áéíóúñü', 'aeiounu'):
        s = s.lower().replace(a, b)
    return re.sub(r'[^a-z0-9]+', '-', s).strip('-')


def archivo(h):
    return f"{h['id']}-{slug(h['alias'])}.md"


ORDEN = ['Sprint 1', 'Sprint 2', 'Sprint 3', 'Sprint 4', 'Release 2 (fuera de v1.0)']
L = ['# Pila de producto de CalisFacu v1.0', '',
     '> Versión en Markdown de `backlog_v2_4sprints.xlsx` (entregable del curso). Si una cambia, la otra se',
     '> actualiza en el mismo commit (SQAP, Sección 10.7). Se regenera con `docs/herramientas/gen_backlog_md.py`.', '',
     '- 4 sprints de **1 semana** · equipo de 2 (Lujan y Galvan) · velocidad supuesta: 20 puntos por semana.',
     '- US-17 a US-20 son historias nuevas propuestas en la revisión del backlog.',
     '- El módulo de pagos (US-13 a US-16) queda **fuera de la v1.0** (Release 2).', '', '## Plan de sprints', '']
for s in ORDEN:
    hs = [h for h in hist if h['sprint'] == s]
    L.append(f"- **{s}** ({sum(h['esfuerzo'] for h in hs)} puntos): "
             + ', '.join(f"[{h['id']} {h['alias']}](../historias/{archivo(h)})" for h in hs))
L += ['', '## Historias', '',
      tabla(['ID', 'Historia', 'Sprint', 'Prioridad', 'Valor', 'Esfuerzo', 'ROI'],
            [[h['id'], h['alias'], h['sprint'], h['prioridad'], h['valor'], h['esfuerzo'], h['roi']] for h in hist])]
L += ['', '## Definition of Ready', ''] + [f'- {a} _(se comprueba en: {b})_' for a, b in dor]
L += ['', '## Definition of Done', ''] + [f'- {a} _(se comprueba en: {b})_' for a, b in dod]
L += ['', '## Hallazgos de la revisión del backlog', '',
      tabla(['N.º', 'Hallazgo', 'Dónde', 'Por qué importa', 'Qué se hizo'], HZ)]
open(f'{DOCS}/backlog/backlog.md', 'w').write('\n'.join(L) + '\n')

# Decisiones del equipo sobre las preguntas abiertas del SRS (Sección 2.3). No están en el Excel del curso.
DECISIONES = [
    '## Decisiones del equipo (13/10/2026)', '',
    'Las preguntas abiertas del SRS (Sección 2.3) se resolvieron así. Si el cliente o el reglamento dicen otra cosa, se cambia',
    'con una solicitud de cambio (SQAP, Sección 16) y se agrega el caso a `casos-calculo.json`.', '',
    tabla(['N.º', 'Pregunta', 'Decisión', 'Dónde se aplica'], [
        ['D1', '¿Se admite 0 en una actividad no presentada?',
         'No. La escala es de 1 a 100; una actividad no presentada se registra con 1.',
         'DTO de la API, `CHECK notas_valor_escala`, web (US-01)'],
        ['D2', '¿Cómo se redondea?',
         'Al entero más cercano, 0,5 hacia arriba (`ROUND_HALF_UP`), solo al final de la suma.',
         '`CalculoService.redondear` (US-02)'],
        ['D3', '¿La nota anual es el promedio de los 3 trimestres?', 'Sí, redondeado con la regla D2.',
         '`CalculoService.notaAnual` (US-02)'],
        ['D4', 'Dimensión sin notas: ¿"pendiente" o 0?',
         'Pendiente: no hay nota trimestral hasta que las cuatro dimensiones tengan al menos una nota.',
         '`CalculoService.notaTrimestral` (US-02)'],
        ['D5', '¿Una licencia cuenta como asistencia?',
         'Sí: porcentaje = (presentes + licencias) / clases dictadas × 100, con un decimal.', 'US-19'],
        ['D6', 'Umbral de asistencia para "en riesgo"',
         'Menos de 80 %, configurable por gestión. Riesgo = promedio < 51 **o** asistencia < 80 %.', 'US-06'],
        ['D7', 'Canal de notificación', 'Aviso dentro del sistema y correo electrónico.', 'US-08, US-12'],
        ['D8', '¿Qué es una "fecha futura" en la asistencia?',
         'Posterior a hoy en Bolivia (America/La_Paz), no a la fecha del servidor (UTC).',
         '`reglas-asistencia.ts` (US-19)'],
        ['D9', '¿Se puede tomar asistencia a una parte del curso?',
         'No: cada clase registra a todos los estudiantes del curso, así "clases dictadas" es igual para todos.',
         '`AsistenciaService.registrar` (US-19)'],
        ['D10', '¿Cómo se sabe que un estudiante ya está inscrito en la gestión?',
         'Por su código RUDE (Registro Único de Estudiantes): no se repite en una gestión; en otra gestión vuelve a inscribirse.',
         '`UNIQUE estudiantes_rude_por_gestion`, `EstudiantesService` (US-18, criterio 3)'],
        ['D11', '¿Qué es "30 minutos sin actividad" en la sesión?',
         'El token dura 30 minutos y la web lo renueva (`POST /api/auth/renovar`) mientras el usuario trabaja; sin peticiones durante 30 minutos, vence.',
         '`AuthService.renovar` (US-17, criterio 4)'],
        ['D12', '¿Qué muestra el historial del estudiante y quién lo ve?',
         'Todas las materias de su curso en cada trimestre, con la nota calculada al consultar. El profesor lo ve si dicta al menos una materia del curso; la secretaria y el administrador, siempre. '
         'Promedio del trimestre = promedio de las notas de las materias, con dos decimales (pendiente si falta alguna); promedio anual = promedio de los trimestres con nota; '
         'tendencia = comparación de los dos últimos trimestres con nota; en riesgo si el promedio anual es menor que 51. Son indicadores, no notas oficiales (las oficiales son D2 y D3 por materia).',
         '`CalculoService.resumenAnual`, `HistorialService` (US-03)'],
    ]), '',
]

R = ['# Reglas de evaluación del Ministerio de Educación', '',
     '> Reglas que implementa el módulo `api/src/modules/calculo` (US-02, RNF-01). Los puntajes de las dimensiones',
     '> se **configuran por gestión** (tabla `dimensiones_evaluacion`, US-18): no se programan fijos.', '',
     tabla(['Regla', 'Valor 2026', 'Fuente', 'Estado'], REG), '', '## Fórmula', '',
     '1. Puntos de una dimensión = promedio de las notas de sus actividades × (puntaje máximo de la dimensión / 100).',
     '2. Nota trimestral = suma de los puntos de las cuatro dimensiones, redondeada al entero (0,5 hacia arriba).',
     '3. Aprueba si la nota es ≥ 51.',
     '4. Si una dimensión todavía no tiene actividades con nota, la nota trimestral queda **pendiente** (D4).', '',
     *DECISIONES,
     '## Casos de ejemplo (semilla de la tabla que firma el cliente, riesgo R02)', '',
     'Son los mismos casos de `api/test/fixtures/casos-calculo.json`, que verifica la prueba `calculo.casos.spec.ts` (M01).', '',
     tabla(['Caso', 'Ser', 'Saber', 'Hacer', 'Autoev.', 'Pts Ser', 'Pts Saber', 'Pts Hacer', 'Pts Autoev.',
            'Sin redondear', 'Nota', 'Resultado'], CAS), '', '## Fuentes', '',
     '- RM 001/2026, Ministerio de Educación (5 de enero de 2026): Ser 10, Saber 45, Hacer 40, Autoevaluación 5.',
     '- Reglamento de Evaluación RM 0190/2024: evaluación trimestral.',
     '- Escala de 1 a 100 y aprobación con 51: vigente desde 2013.',
     '- Verificar contra el texto oficial antes de firmar la tabla.']
open(f'{DOCS}/backlog/reglas-ministerio.md', 'w').write('\n'.join(R) + '\n')

for h in hist:
    L = [f"# {h['id']} · {h['alias']}", '',
         f"- **Sprint:** {h['sprint']} · **Prioridad:** {h['prioridad']} · **Puntos:** valor {h['valor']}, "
         f"esfuerzo {h['esfuerzo']} · **Estado:** {h['estado']}", '',
         '## Historia', '', h['enunciado'], '', '## Notas y dependencias', '', h['comentario'], '',
         '## Criterios de aceptación', '']
    for c in CA.get(h['id'], []):
        extra = ' · '.join(x for x in [c['prueba'] and f"Prueba: {c['prueba']}", c['vinculo'] and f"SQAP: {c['vinculo']}"] if x)
        L.append(f"{c['n']}. {c['texto']}" + (f"  \n   _{extra}_" if extra else ''))
    if T.get(h['id']):
        L += ['', '## Tareas', '', tabla(['ID', 'Tarea', 'Tipo', 'Responsable', 'Revisor', 'Horas', 'Cubre'],
                                        [[t['id'], t['tarea'], t['tipo'], t['resp'], t['rev'], t['horas'], t['vinc']] for t in T[h['id']]])]
    L += ['', '## Antes de empezar', '',
          '- [ ] Cumple la Definition of Ready (`docs/backlog/backlog.md`).',
          '- [ ] No quedan ❓ sin respuesta del cliente.',
          '- [ ] Se trabaja directo en `develop`, con commits que citan la historia (ADR-0008).']
    open(f"{DOCS}/historias/{archivo(h)}", 'w').write('\n'.join(L) + '\n')

tr = sorted([t for k, v in T.items() if k in ('Calidad', 'Infra') for t in v], key=lambda t: t['id'])
open(f'{DOCS}/backlog/tareas-transversales.md', 'w').write('\n'.join([
    '# Tareas transversales (calidad e infraestructura)', '',
    tabla(['ID', 'Sprint', 'Tarea', 'Tipo', 'Responsable', 'Revisor', 'Horas', 'Cubre'],
          [[t['id'], t['sprint'], t['tarea'], t['tipo'], t['resp'], t['rev'], t['horas'], t['vinc']] for t in tr])]) + '\n')

casos = [dict(caso=c[0], promedios=dict(SER=c[1], SABER=c[2], HACER=c[3], AUTOEVALUACION=c[4]),
              sinRedondear=c[9], nota=c[10], resultado=c[11]) for c in CAS]
json.dump(casos, open(RAIZ / 'api/test/fixtures/casos-calculo.json', 'w'), ensure_ascii=False, indent=2)
print('historias', len(hist))
