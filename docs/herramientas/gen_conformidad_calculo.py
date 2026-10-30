"""Genera el acta de conformidad del cálculo (T-004, riesgo R02) en Markdown y Word.

Lee los casos de api/test/fixtures/casos-calculo.json (los que verifica calculo.casos.spec.ts en cada build)
y las decisiones D1 a D9 de docs/backlog/reglas-ministerio.md, así lo que firma el cliente es lo mismo que prueba
el pipeline. El PDF se obtiene del .docx con:
  soffice --headless --convert-to pdf --outdir docs/actas docs/actas/conformidad-reglas-calculo.docx
"""
import json
import pathlib
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = pathlib.Path(__file__).resolve().parents[2]
SALIDA = RAIZ / 'docs/actas/conformidad-reglas-calculo'
CASOS = json.loads((RAIZ / 'api/test/fixtures/casos-calculo.json').read_text(encoding='utf-8'))
REGLAS = (RAIZ / 'docs/backlog/reglas-ministerio.md').read_text(encoding='utf-8')
DECISIONES = [
    [c.strip() for c in fila.strip().strip('|').split('|')[:3]]
    for fila in REGLAS.splitlines()
    if re.match(r'^\|\s*D\d+\s*\|', fila)
]
PUNTAJES = {'SER': 10, 'SABER': 45, 'HACER': 40, 'AUTOEVALUACION': 5}  # RM 001/2026, igual que el seed de 2026
DIMS = ['SER', 'SABER', 'HACER', 'AUTOEVALUACION']


def num(x):
    texto = f'{x:.2f}'.rstrip('0').rstrip('.')
    return texto.replace('.', ',')


def puntos(caso):
    return [num(caso['promedios'][d] * PUNTAJES[d] / 100) for d in DIMS]


def sin_md(texto):
    return re.sub(r'\*\*|`', '', texto)


ANUALES = [  # calculo.service.spec.ts, describe notaAnual
    ([50, 51, 51], '50,67', 51, 'Aprueba'),
    ([50, 50, 51], '50,33', 50, 'Reprueba'),
]
PENDIENTE = ('Ser 80 · Saber 70 · Hacer 80 · Autoevaluación sin notas', 'Pendiente (falta Autoevaluación)')

ENCABEZADO = [
    ('Proyecto', 'CalisFacu v1.0 · Equipo L-001 · COM450 Calidad de Software'),
    ('Cliente', 'Unidad Educativa Nazareno'),
    ('Tarea y riesgo', 'T-004 · R02 (cálculo distinto al reglamento) · US-02'),
    ('Reglas', 'docs/backlog/reglas-ministerio.md (RM 001/2026, decisiones del 13/10/2026)'),
    ('Pruebas que lo verifican', 'calculo.casos.spec.ts y calculo.service.spec.ts, en cada build de Jenkins (M01)'),
    ('Fecha y lugar de la firma', ''),
]
PROPOSITO = (
    'El cliente confirma que CalisFacu calcula la nota trimestral y la nota anual como lo hace la unidad educativa '
    'según la RM 001/2026. Cada caso de esta acta es una prueba automática: si el cliente marca un caso o una '
    'decisión como "Cambiar", el equipo abre una solicitud de cambio (SQAP, Sección 16), corrige la prueba y el '
    'código, y vuelve a presentar el acta.'
)
FORMULA = [
    'Puntos de una dimensión = promedio de las notas de sus actividades × (puntaje máximo de la dimensión ÷ 100).',
    'Nota trimestral = suma de los puntos de las cuatro dimensiones, redondeada al entero solo al final '
    '(0,5 hacia arriba).',
    'Aprueba si la nota es 51 o más. Escala de cada actividad: 1 a 100.',
]
FIRMAS = [
    ('Fatima Carreño', 'Cliente: profesora representante'),
    ('Fabiana Parrado', 'Cliente: directora'),
    ('Lujan', 'Autor (Líder y Arquitecto)'),
    ('Galvan', 'Responsable SQA'),
]
TITULO = 'Acta de conformidad de las reglas de cálculo'


def markdown():
    l = [f'# {TITULO}', '', '> Generada con `python3 docs/herramientas/gen_conformidad_calculo.py`. La versión para imprimir y',
         '> firmar es `conformidad-reglas-calculo.pdf`; el acta firmada se escanea y se guarda junto a ella.', '']
    l += [f'- **{k}:** {v or "____________________"}' for k, v in ENCABEZADO]
    l += ['', PROPOSITO, '', '## 1. Fórmula y puntajes (RM 001/2026)', '']
    l += [f'{i}. {f}' for i, f in enumerate(FORMULA, 1)]
    l += ['', '| Dimensión | Ser | Saber | Hacer | Autoevaluación | Total |', '| --- | --- | --- | --- | --- | --- |',
          f'| Puntaje máximo | {" | ".join(str(PUNTAJES[d]) for d in DIMS)} | 100 |', '',
          '## 2. Casos de la nota trimestral', '',
          'Promedio de cada dimensión (escala 1 a 100), puntos que aporta y nota final.', '',
          '| N.º | Caso | Ser | Saber | Hacer | Autoev. | Puntos (Ser + Saber + Hacer + Autoev.) | Sin redondear | Nota | Resultado | Conforme |',
          '| --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- |']
    for i, c in enumerate(CASOS, 1):
        p = c['promedios']
        l.append(f'| {i} | {c["caso"]} | {" | ".join(str(p[d]) for d in DIMS)} | {" + ".join(puntos(c))} | '
                 f'{num(c["sinRedondear"])} | {c["nota"]} | {c["resultado"]} | ☐ Sí ☐ No |')
    l += ['', '## 3. Nota anual y nota pendiente', '',
          '| N.º | Caso | Datos | Cálculo | Nota | Resultado | Conforme |', '| --- | --- | --- | --- | --- | --- | --- |']
    n = len(CASOS)
    for j, (t, prom, nota, res) in enumerate(ANUALES, n + 1):
        l.append(f'| {j} | Nota anual (D3) | Trimestres {", ".join(map(str, t))} | ({" + ".join(map(str, t))}) ÷ 3 = {prom} | '
                 f'{nota} | {res} | ☐ Sí ☐ No |')
    l.append(f'| {n + 3} | Dimensión sin notas (D4) | {PENDIENTE[0]} | No se calcula | — | {PENDIENTE[1]} | ☐ Sí ☐ No |')
    l += ['', '## 4. Decisiones del equipo', '',
          'Las tomó el equipo porque el reglamento no las define o el SRS las dejó abiertas. El cliente marca cada una.', '',
          '| N.º | Pregunta | Decisión | De acuerdo | Cambiar (observación) |', '| --- | --- | --- | --- | --- |']
    l += [f'| {d} | {p} | {sin_md(dec)} | ☐ | |' for d, p, dec in DECISIONES]
    l += ['', '## 5. Observaciones del cliente', '', '_____________________________________________', '',
          '_____________________________________________', '', '## 6. Decisión', '',
          '- [ ] Conforme: las reglas y los casos quedan aprobados.',
          '- [ ] Conforme con cambios: se abren las solicitudes de cambio anotadas en la Sección 5.',
          '- [ ] No conforme: se corrige y se vuelve a presentar el ____/____/2026.', '', '## 7. Firmas', '',
          '| Nombre | Rol | Firma |', '| --- | --- | --- |']
    l += [f'| {n or ""} | {r} | |' for n, r in FIRMAS]
    return '\n'.join(l) + '\n'


def sombrear(celda, color):
    tc = celda._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), color)
    tc.append(shd)


def tabla(doc, filas, anchos, tam=9, encabezado=True):
    t = doc.add_table(rows=len(filas), cols=len(filas[0]))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, ancho in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), anchos):
        col.set(qn('w:w'), str(int(Cm(ancho).twips)))
    for fila in t.rows:
        fila_pr = fila._tr.get_or_add_trPr()
        fila_pr.append(OxmlElement('w:cantSplit'))
    encabezado_pr = t.rows[0]._tr.get_or_add_trPr()
    if encabezado:
        encabezado_pr.append(OxmlElement('w:tblHeader'))
    for i, fila in enumerate(filas):
        for j, valor in enumerate(fila):
            celda = t.cell(i, j)
            celda.width = Cm(anchos[j])
            p = celda.paragraphs[0]
            r = p.add_run(str(valor))
            r.font.size = Pt(tam)
            if encabezado and i == 0:
                r.bold = True
                sombrear(celda, 'DEEBF7')
    doc.add_paragraph()
    return t


def titulo(doc, texto):
    h = doc.add_heading(texto, level=2)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)


def word():
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(27.94), Cm(21.59)
    for lado in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, lado, Cm(1.6))
    estilo = doc.styles['Normal']
    estilo.font.name = 'Calibri'
    estilo.font.size = Pt(10.5)
    pie = sec.footer.paragraphs[0]
    pie.text = 'CalisFacu v1.0 · Acta de conformidad de las reglas de cálculo (T-004) · Equipo L-001'
    pie.runs[0].font.size = Pt(8)

    h = doc.add_heading(TITULO, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tabla(doc, [[k, v] for k, v in ENCABEZADO], [6, 18.5], tam=10, encabezado=False)
    doc.add_paragraph(PROPOSITO)

    titulo(doc, '1. Fórmula y puntajes (RM 001/2026)')
    for f in FORMULA:
        doc.add_paragraph(f, style='List Number')
    tabla(doc, [['Dimensión', 'Ser', 'Saber', 'Hacer', 'Autoevaluación', 'Total'],
                ['Puntaje máximo', *[PUNTAJES[d] for d in DIMS], 100]], [5, 3, 3, 3, 4, 3])

    titulo(doc, '2. Casos de la nota trimestral')
    doc.add_paragraph('Promedio de cada dimensión (escala 1 a 100), puntos que aporta y nota final.')
    filas = [['N.º', 'Caso', 'Ser', 'Saber', 'Hacer', 'Autoev.', 'Puntos (Ser + Saber + Hacer + Autoev.)',
              'Sin redondear', 'Nota', 'Resultado', 'Conforme']]
    for i, c in enumerate(CASOS, 1):
        p = c['promedios']
        filas.append([i, c['caso'], *[p[d] for d in DIMS], ' + '.join(puntos(c)), num(c['sinRedondear']),
                      c['nota'], c['resultado'], '☐ Sí  ☐ No'])
    tabla(doc, filas, [1, 5.7, 1.3, 1.3, 1.3, 1.7, 4.6, 2, 1.3, 2, 2.4])

    titulo(doc, '3. Nota anual y nota pendiente')
    n = len(CASOS)
    filas = [['N.º', 'Caso', 'Datos', 'Cálculo', 'Nota', 'Resultado', 'Conforme']]
    for j, (t, prom, nota, res) in enumerate(ANUALES, n + 1):
        filas.append([j, 'Nota anual (D3)', 'Trimestres ' + ', '.join(map(str, t)),
                      f'({" + ".join(map(str, t))}) ÷ 3 = {prom}', nota, res, '☐ Sí  ☐ No'])
    filas.append([n + 3, 'Dimensión sin notas (D4)', PENDIENTE[0], 'No se calcula', '—', PENDIENTE[1], '☐ Sí  ☐ No'])
    tabla(doc, filas, [1, 4, 6.5, 5, 1.3, 4.4, 2.4])

    titulo(doc, '4. Decisiones del equipo')
    doc.add_paragraph('Las tomó el equipo porque el reglamento no las define o el SRS las dejó abiertas. '
                      'El cliente marca cada una.')
    tabla(doc, [['N.º', 'Pregunta', 'Decisión', 'De acuerdo', 'Cambiar (observación)']] +
          [[d, p, sin_md(dec), '☐', ''] for d, p, dec in DECISIONES], [1, 6, 10, 2, 5.5])

    titulo(doc, '5. Observaciones del cliente')
    for _ in range(4):
        doc.add_paragraph('_' * 120)

    titulo(doc, '6. Decisión')
    for texto in ('☐ Conforme: las reglas y los casos quedan aprobados.',
                  '☐ Conforme con cambios: se abren las solicitudes de cambio anotadas en la Sección 5.',
                  '☐ No conforme: se corrige y se vuelve a presentar el ____/____/2026.'):
        doc.add_paragraph(texto)

    titulo(doc, '7. Firmas')
    t = tabla(doc, [['Nombre', 'Rol', 'Firma']] + [[n, r, ''] for n, r in FIRMAS], [7, 8, 9.5], tam=10)
    for fila in t.rows[1:]:
        fila.height = Cm(1.6)
    doc.save(str(SALIDA) + '.docx')


if __name__ == '__main__':
    pathlib.Path(str(SALIDA) + '.md').write_text(markdown(), encoding='utf-8')
    word()
    print(f'{SALIDA}.md y .docx')
